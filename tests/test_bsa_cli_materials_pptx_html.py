"""Tests for v1.4.5 bsa materials pptx + html support
(closes lifecycle review rec #3 first half).

New extractors:
  _convert_pptx — python-pptx lazy-imported; H2 per visible slide,
                  text from each shape, speaker notes as H3.
  _convert_html — stdlib html.parser + html.unescape; analyst-grep-
                  friendly markdown (headings, paragraphs, lists,
                  links, code).

Coverage:
  - extension map: .pptx → "pptx"; .html / .htm → "html".
  - pptx: happy path (title + body), multi-slide order, hidden-skip,
    speaker notes, table rendering, ConversionUnavailable when dep
    missing, empty-deck fallback message.
  - html: happy path, script/style stripped, title-as-h1 fallback,
    title NOT used when h1 present, lists, pre/code preservation,
    encoding fallback (latin-1), entity decoding, href rendering,
    nested inline emphasis.
  - CLI: dry-run preview + commit for both formats.
"""

from __future__ import annotations

import importlib
import re
import subprocess
import sys
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parent.parent


def _make_workspace(tmp_path: Path) -> Path:
    ws = tmp_path / "ws"
    (ws / "analysis" / "proposals" / "stage1" / "inputs").mkdir(parents=True)
    return ws


def _run_cli(*args: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, str(REPO_ROOT / "scripts" / "bsa_cli.py"), *args],
        capture_output=True, text=True, check=False,
    )


def _ensure_pptx_available() -> None:
    """Skip if python-pptx isn't installed in the test environment.
    Mirrors the pdf/docx test pattern; CI is expected to have it but
    a barebones developer venv may not."""
    try:
        importlib.import_module("pptx")
    except ImportError:
        pytest.skip("python-pptx not installed")


# ---- extension map --------------------------------------------------


def test_extension_map_includes_pptx_and_html() -> None:
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import _EXT_TO_KIND
    finally:
        sys.path.pop(0)
    assert _EXT_TO_KIND[".pptx"] == "pptx"
    assert _EXT_TO_KIND[".html"] == "html"
    assert _EXT_TO_KIND[".htm"] == "html"


# ---- _convert_pptx --------------------------------------------------


def _make_pptx(path: Path, slides: list[dict]) -> None:
    """Build a minimal pptx at `path`. Each slides[i] is a dict with
    optional keys: title (str), body (list[str]), notes (str),
    hidden (bool), table (list[list[str]])."""
    _ensure_pptx_available()
    from pptx import Presentation
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]  # truly blank
    title_layout = prs.slide_layouts[5]  # title-only
    for spec in slides:
        layout = title_layout if spec.get("title") is not None else blank_layout
        slide = prs.slides.add_slide(layout)
        if spec.get("title") is not None:
            slide.shapes.title.text = spec["title"]
        for body_text in spec.get("body", []):
            from pptx.util import Inches
            tb = slide.shapes.add_textbox(
                Inches(1), Inches(2), Inches(6), Inches(1)
            )
            tf = tb.text_frame
            tf.text = body_text
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec["notes"]
        if spec.get("hidden"):
            slide.element.set("show", "0")
        if spec.get("table"):
            from pptx.util import Inches
            rows = spec["table"]
            table_shape = slide.shapes.add_table(
                rows=len(rows), cols=len(rows[0]),
                left=Inches(1), top=Inches(3),
                width=Inches(6), height=Inches(2),
            )
            for r_idx, row in enumerate(rows):
                for c_idx, val in enumerate(row):
                    table_shape.table.cell(r_idx, c_idx).text = val
    prs.save(str(path))


def test_convert_pptx_happy_path(tmp_path: Path) -> None:
    _ensure_pptx_available()
    from scripts._bsa_cli_materials import _convert_pptx
    p = tmp_path / "deck.pptx"
    _make_pptx(p, [
        {"title": "Intro", "body": ["bullet one", "bullet two"]},
    ])
    out = _convert_pptx(p)
    assert "## Slide 1: Intro" in out
    assert "bullet one" in out
    assert "bullet two" in out


def test_convert_pptx_multi_slide_preserves_order(tmp_path: Path) -> None:
    _ensure_pptx_available()
    from scripts._bsa_cli_materials import _convert_pptx
    p = tmp_path / "deck.pptx"
    _make_pptx(p, [
        {"title": "Alpha"},
        {"title": "Bravo"},
        {"title": "Charlie"},
    ])
    out = _convert_pptx(p)
    pos_a = out.index("Alpha")
    pos_b = out.index("Bravo")
    pos_c = out.index("Charlie")
    assert pos_a < pos_b < pos_c


def test_convert_pptx_skips_hidden_slides(tmp_path: Path) -> None:
    _ensure_pptx_available()
    from scripts._bsa_cli_materials import _convert_pptx
    p = tmp_path / "deck.pptx"
    _make_pptx(p, [
        {"title": "Visible1"},
        {"title": "HiddenSlide", "hidden": True},
        {"title": "Visible2"},
    ])
    out = _convert_pptx(p)
    assert "Visible1" in out
    assert "Visible2" in out
    assert "HiddenSlide" not in out
    # Slide numbering reflects VISIBLE order (1, 2 — not 1, 3).
    assert "## Slide 1:" in out
    assert "## Slide 2:" in out
    assert "## Slide 3:" not in out


def test_convert_pptx_includes_speaker_notes(tmp_path: Path) -> None:
    _ensure_pptx_available()
    from scripts._bsa_cli_materials import _convert_pptx
    p = tmp_path / "deck.pptx"
    _make_pptx(p, [
        {"title": "S1", "notes": "this is the speaker note text"},
    ])
    out = _convert_pptx(p)
    assert "### Speaker notes" in out
    assert "this is the speaker note text" in out


def test_convert_pptx_renders_table_as_pipe_rows(tmp_path: Path) -> None:
    _ensure_pptx_available()
    from scripts._bsa_cli_materials import _convert_pptx
    p = tmp_path / "deck.pptx"
    _make_pptx(p, [
        {
            "title": "TableSlide",
            "table": [["H1", "H2"], ["a1", "a2"], ["b1", "b2"]],
        },
    ])
    out = _convert_pptx(p)
    assert "H1 | H2" in out
    assert "a1 | a2" in out
    assert "b1 | b2" in out


def test_convert_pptx_empty_deck_yields_placeholder(tmp_path: Path) -> None:
    _ensure_pptx_available()
    from scripts._bsa_cli_materials import _convert_pptx
    p = tmp_path / "deck.pptx"
    _make_pptx(p, [])  # no slides
    out = _convert_pptx(p)
    assert "no visible slides" in out


def test_convert_pptx_unavailable_when_dep_missing(tmp_path: Path) -> None:
    """Mock the import to simulate missing python-pptx."""
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import (
            ConversionUnavailable,
            _convert_pptx,
        )
    finally:
        sys.path.pop(0)
    p = tmp_path / "fake.pptx"
    p.write_text("not a real pptx", encoding="utf-8")
    # Sabotage the import so the lazy `from pptx import ...` raises.
    saved = sys.modules.pop("pptx", None)
    sys.modules["pptx"] = None  # type: ignore[assignment]
    try:
        with pytest.raises(ConversionUnavailable):
            _convert_pptx(p)
    finally:
        if saved is not None:
            sys.modules["pptx"] = saved
        else:
            sys.modules.pop("pptx", None)


# ---- _convert_html --------------------------------------------------


def test_convert_html_happy_path(tmp_path: Path) -> None:
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import _convert_html
    finally:
        sys.path.pop(0)
    p = tmp_path / "page.html"
    p.write_text(
        "<html><head><title>Doc</title></head><body>"
        "<h1>Top</h1><p>First para.</p><p>Second.</p>"
        "</body></html>",
        encoding="utf-8",
    )
    out = _convert_html(p)
    assert "# Top" in out
    assert "First para." in out
    assert "Second." in out


def test_convert_html_strips_script_and_style(tmp_path: Path) -> None:
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import _convert_html
    finally:
        sys.path.pop(0)
    p = tmp_path / "page.html"
    p.write_text(
        "<html><head>"
        "<script>console.log('SECRET-JS-PAYLOAD')</script>"
        "<style>.x{color: red /* SECRET-CSS */}</style>"
        "</head><body><p>Visible body</p></body></html>",
        encoding="utf-8",
    )
    out = _convert_html(p)
    assert "Visible body" in out
    assert "SECRET-JS-PAYLOAD" not in out
    assert "SECRET-CSS" not in out


def test_convert_html_title_used_when_no_h1(tmp_path: Path) -> None:
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import _convert_html
    finally:
        sys.path.pop(0)
    p = tmp_path / "page.html"
    p.write_text(
        "<html><head><title>The Title</title></head>"
        "<body><p>body para no h1</p></body></html>",
        encoding="utf-8",
    )
    out = _convert_html(p)
    assert out.startswith("# The Title"), (
        f"Title must become H1 when body has no H1; got: {out!r}"
    )


def test_convert_html_title_not_used_when_h1_present(tmp_path: Path) -> None:
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import _convert_html
    finally:
        sys.path.pop(0)
    p = tmp_path / "page.html"
    p.write_text(
        "<html><head><title>FromTitle</title></head>"
        "<body><h1>FromBody</h1><p>x</p></body></html>",
        encoding="utf-8",
    )
    out = _convert_html(p)
    assert "FromBody" in out
    assert "FromTitle" not in out


def test_convert_html_lists_render_as_bullets(tmp_path: Path) -> None:
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import _convert_html
    finally:
        sys.path.pop(0)
    p = tmp_path / "page.html"
    p.write_text(
        "<html><body><ul><li>alpha</li><li>beta</li></ul></body></html>",
        encoding="utf-8",
    )
    out = _convert_html(p)
    assert "- alpha" in out
    assert "- beta" in out


def test_convert_html_preserves_pre_block(tmp_path: Path) -> None:
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import _convert_html
    finally:
        sys.path.pop(0)
    p = tmp_path / "page.html"
    p.write_text(
        "<html><body><pre>line1\n  indent line2</pre></body></html>",
        encoding="utf-8",
    )
    out = _convert_html(p)
    assert "```" in out
    assert "line1" in out
    assert "  indent line2" in out  # whitespace inside <pre> preserved


def test_convert_html_renders_inline_code(tmp_path: Path) -> None:
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import _convert_html
    finally:
        sys.path.pop(0)
    p = tmp_path / "page.html"
    p.write_text(
        "<html><body><p>Run <code>npm test</code>.</p></body></html>",
        encoding="utf-8",
    )
    out = _convert_html(p)
    assert "`npm test`" in out


def test_convert_html_link_renders_as_markdown(tmp_path: Path) -> None:
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import _convert_html
    finally:
        sys.path.pop(0)
    p = tmp_path / "page.html"
    p.write_text(
        "<html><body><p>See <a href=\"https://x.example/y\">the docs</a>.</p></body></html>",
        encoding="utf-8",
    )
    out = _convert_html(p)
    assert "[the docs](https://x.example/y)" in out


def test_convert_html_decodes_entities(tmp_path: Path) -> None:
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import _convert_html
    finally:
        sys.path.pop(0)
    p = tmp_path / "page.html"
    p.write_text(
        "<html><body><p>A &amp; B &lt; C &gt; D &copy; 2026</p></body></html>",
        encoding="utf-8",
    )
    out = _convert_html(p)
    assert "A & B" in out
    assert "< C >" in out or "&lt;" not in out
    assert "©" in out or "&copy;" not in out


def test_convert_html_encoding_fallback_latin1(tmp_path: Path) -> None:
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import _convert_html
    finally:
        sys.path.pop(0)
    p = tmp_path / "page.html"
    # Latin-1 byte 0xe9 = é. Invalid UTF-8 standalone.
    p.write_bytes(
        b"<html><body><p>Caf\xe9 chic</p></body></html>",
    )
    out = _convert_html(p)
    assert "Café chic" in out


def test_convert_html_emphasis(tmp_path: Path) -> None:
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import _convert_html
    finally:
        sys.path.pop(0)
    p = tmp_path / "page.html"
    p.write_text(
        "<html><body><p>Be <strong>bold</strong> and <em>italic</em>.</p></body></html>",
        encoding="utf-8",
    )
    out = _convert_html(p)
    assert "**bold**" in out
    assert "*italic*" in out


# ---- CLI integration ------------------------------------------------


def test_cli_dry_run_lists_pptx_and_html(tmp_path: Path) -> None:
    _ensure_pptx_available()
    ws = _make_workspace(tmp_path)
    src = tmp_path / "src"
    src.mkdir()
    _make_pptx(src / "deck.pptx", [{"title": "S1", "body": ["body"]}])
    (src / "page.html").write_text(
        "<html><body><h1>Hi</h1></body></html>", encoding="utf-8",
    )
    res = _run_cli(f"--workspace={ws}", "materials", str(src))
    assert res.returncode == 0, res.stderr
    assert "PPTX: 1" in res.stdout
    assert "HTML: 1" in res.stdout


def test_cli_commit_pptx_and_html(tmp_path: Path) -> None:
    _ensure_pptx_available()
    ws = _make_workspace(tmp_path)
    src = tmp_path / "src"
    src.mkdir()
    _make_pptx(src / "deck.pptx", [{"title": "Slide", "body": ["text"]}])
    (src / "page.html").write_text(
        "<html><body><h1>Heading</h1><p>Body.</p></body></html>",
        encoding="utf-8",
    )
    res = _run_cli(
        f"--workspace={ws}", "materials", str(src), "--commit",
    )
    assert res.returncode == 0, res.stderr
    inputs = ws / "analysis" / "proposals" / "stage1" / "inputs"
    pptx_staged = next(inputs.glob("source_*_deck.md"))
    html_staged = next(inputs.glob("source_*_page.md"))
    assert "## Slide 1: Slide" in pptx_staged.read_text(encoding="utf-8")
    assert "# Heading" in html_staged.read_text(encoding="utf-8")


def test_cli_unsupported_htm_alias(tmp_path: Path) -> None:
    """`.htm` (3-letter) variant must be classified the same as `.html`."""
    ws = _make_workspace(tmp_path)
    src = tmp_path / "src"
    src.mkdir()
    (src / "legacy.htm").write_text(
        "<html><body><p>hi</p></body></html>", encoding="utf-8",
    )
    res = _run_cli(f"--workspace={ws}", "materials", str(src), "--commit")
    assert res.returncode == 0, res.stderr
    inputs = ws / "analysis" / "proposals" / "stage1" / "inputs"
    staged = next(inputs.glob("source_*_legacy.md"))
    assert "hi" in staged.read_text(encoding="utf-8")


# ---- v1.4.5 R1 fixes — Codex review round 1 -------------------------


def test_pptx_skips_show_false_attribute(tmp_path: Path) -> None:
    """R1 MAJOR #1: OOXML booleans accept `0`/`1`/`false`/`true`.
    Hidden slides authored with `show="false"` (not just `show="0"`)
    must also be skipped."""
    _ensure_pptx_available()
    from scripts._bsa_cli_materials import _convert_pptx
    p = tmp_path / "deck.pptx"
    _make_pptx(p, [
        {"title": "VisibleSlide"},
        {"title": "HiddenFalse"},  # we'll set show="false" below
        {"title": "VisibleEnd"},
    ])
    # Re-open and patch the middle slide's show attribute.
    from pptx import Presentation
    prs = Presentation(str(p))
    prs.slides[1].element.set("show", "false")
    prs.save(str(p))
    out = _convert_pptx(p)
    assert "VisibleSlide" in out
    assert "VisibleEnd" in out
    assert "HiddenFalse" not in out


def test_html_respects_meta_charset_cp1252(tmp_path: Path) -> None:
    """R1 MAJOR #2: an HTML page whose <meta charset> declares
    Windows-1252 (legacy web exports) must decode correctly. The
    pre-fix latin-1 fallback turned cp1252-specific bytes (smart
    quotes, em-dash, copyright) into control characters."""
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import _convert_html
    finally:
        sys.path.pop(0)
    p = tmp_path / "page.html"
    # 0x93 = U+201C "left double quotation mark" in cp1252 (control
    # in latin-1). 0x97 = U+2014 "em dash". 0xA9 = "©" in both.
    body = (
        b"<html><head>"
        b"<meta charset=\"windows-1252\">"
        b"<title>cp1252 page</title>"
        b"</head><body><p>Smart \x93quote\x94 \x97 dash \xa9 2026</p></body></html>"
    )
    p.write_bytes(body)
    out = _convert_html(p)
    assert "“quote”" in out, (
        f"cp1252 smart quotes must decode; got: {out!r}"
    )
    assert "—" in out  # em-dash
    assert "©" in out


def test_html_respects_utf8_bom(tmp_path: Path) -> None:
    """R1 MAJOR #2: a UTF-8 BOM at the file start must be stripped
    AND the file must decode as utf-8 (not fall through to latin-1)."""
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import _convert_html
    finally:
        sys.path.pop(0)
    p = tmp_path / "page.html"
    body = b"\xef\xbb\xbf" + (
        "<html><body><p>café — résumé</p></body></html>".encode("utf-8")
    )
    p.write_bytes(body)
    out = _convert_html(p)
    assert "café" in out
    assert "résumé" in out


def test_restage_pptx_preserves_document_sourcetype(tmp_path: Path) -> None:
    """R1 MAJOR #3: when --restage-changed updates a PPTX row, its
    SourceType must remain `document`, not silently demote to
    `process_note`."""
    _ensure_pptx_available()
    import csv as _csv
    ws = _make_workspace(tmp_path)
    src = tmp_path / "src"
    src.mkdir()
    pptx_path = src / "deck.pptx"
    _make_pptx(pptx_path, [{"title": "v1", "body": ["body v1"]}])
    # First stage.
    res1 = _run_cli(f"--workspace={ws}", "materials", str(src), "--commit")
    assert res1.returncode == 0
    mp = ws / "analysis" / "proposals" / "stage1" / "source_manifest.csv"
    with mp.open(newline="") as fh:
        rows = list(_csv.DictReader(fh))
    assert len(rows) == 1
    assert rows[0]["SourceType"] == "document"
    # Mutate the deck so restage triggers (overwrite with different content).
    _make_pptx(pptx_path, [{"title": "v2 changed", "body": ["body v2"]}])
    res2 = _run_cli(
        f"--workspace={ws}", "materials", str(src),
        "--commit", "--restage-changed",
    )
    assert res2.returncode == 0, res2.stderr
    with mp.open(newline="") as fh:
        rows_after = list(_csv.DictReader(fh))
    assert rows_after[0]["SourceType"] == "document", (
        f"PPTX restage must keep SourceType=document; got: "
        f"{rows_after[0]['SourceType']!r}"
    )


def test_html_heading_inside_link_renders_separately(tmp_path: Path) -> None:
    """R1 MINOR #4: `<a href="x"><h1>title</h1></a>` should NOT
    produce an empty heading marker followed by a stray link.
    Either the link wraps text (correct) or the heading lands at
    body level. The fix flushes the link before the block tag opens."""
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import _convert_html
    finally:
        sys.path.pop(0)
    p = tmp_path / "page.html"
    p.write_text(
        "<html><body>"
        "<a href=\"https://x.example\"><h1>Wrapped Heading</h1></a>"
        "</body></html>",
        encoding="utf-8",
    )
    out = _convert_html(p)
    # No empty heading marker (a `# ` followed only by whitespace
    # or another marker).
    assert not re.search(r"^# *\n", out, flags=re.MULTILINE), (
        f"empty heading marker present; got: {out!r}"
    )
    # Heading text must appear under a `# ` marker.
    assert re.search(r"^# +Wrapped Heading", out, flags=re.MULTILINE), (
        f"heading text lost; got: {out!r}"
    )


def test_html_title_with_inline_tags_normalized(tmp_path: Path) -> None:
    """R1 MINOR #5: `<title>A <b>B</b> C</title>` must produce
    title 'A B C' (whitespace preserved, no stray ** markers)."""
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import _convert_html
    finally:
        sys.path.pop(0)
    p = tmp_path / "page.html"
    p.write_text(
        "<html><head><title>A <b>B</b> C</title></head>"
        "<body><p>x</p></body></html>",
        encoding="utf-8",
    )
    out = _convert_html(p)
    assert "# A B C" in out, f"title fragments lost spaces; got: {out!r}"
    assert "**" not in out, (
        f"stray ** markers from title leaked into body; got: {out!r}"
    )


def test_html_heading_inside_link_no_stray_href(tmp_path: Path) -> None:
    """R2 NEW MINOR fix: when `<h1>` opens inside `<a>`, the pre-
    flush of the (empty) link buffer must NOT emit a bare `<href>`
    before the heading marker. The heading text lands at body level
    under the `# `, and the link is dropped."""
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import _convert_html
    finally:
        sys.path.pop(0)
    p = tmp_path / "page.html"
    p.write_text(
        "<html><body>"
        "<a href=\"https://x.example/page\"><h1>Wrapped</h1></a>"
        "</body></html>",
        encoding="utf-8",
    )
    out = _convert_html(p)
    assert "<https://x.example/page>" not in out, (
        f"empty link buf must NOT emit stray <href> before heading; "
        f"got: {out!r}"
    )
    assert "# Wrapped" in out


def test_html_link_with_no_text_but_href_still_renders_on_close(
    tmp_path: Path,
) -> None:
    """R2 NEW MINOR fix sanity check: the legacy behavior for an
    empty-text link CLOSED via `</a>` (not pre-flushed by a block
    tag) must still emit `<href>` so naked anchors stay traceable."""
    sys.path.insert(0, str(REPO_ROOT))
    try:
        from scripts._bsa_cli_materials import _convert_html
    finally:
        sys.path.pop(0)
    p = tmp_path / "page.html"
    p.write_text(
        "<html><body><p>Naked: "
        "<a href=\"https://y.example\"></a></p></body></html>",
        encoding="utf-8",
    )
    out = _convert_html(p)
    assert "<https://y.example>" in out, (
        f"href-only link via </a> must still emit <href>; got: {out!r}"
    )


def test_pptx_table_preserves_intra_cell_newlines(tmp_path: Path) -> None:
    """R1 MINOR #6: cell text containing `\\n` must be preserved as
    `<br>` (not flattened to space)."""
    _ensure_pptx_available()
    from scripts._bsa_cli_materials import _convert_pptx
    p = tmp_path / "deck.pptx"
    _make_pptx(p, [
        {
            "title": "Multi-line cell",
            "table": [
                ["H1", "H2"],
                ["line1\nline2", "single"],
            ],
        },
    ])
    out = _convert_pptx(p)
    assert "line1<br>line2" in out, (
        f"intra-cell newline must become <br>; got snippet: {out!r}"
    )


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-v"]))
